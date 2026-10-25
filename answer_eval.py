"""回答層級評估：對每一題實際產生回答，再用 LLM 依照 ground_truth 評分。

評分項目：
  正確性（0/1/2）：2 = 涵蓋標準答案的所有關鍵事實；1 = 部分正確或缺少關鍵數字、條件；0 = 錯誤、矛盾或回答不知道
  忠實：回答的每個具體主張都能在送出的條文中找到依據；外部法規只說明依據、沒有自行補充內容
"""
import json
import re
from datetime import datetime

import pandas as pd

from rag import (MAX_REF_ADD, TOP_K_RERANKED, TOP_K_RETRIEVED, build_context,
                 claude_client, format_chunk, generate, response_text)
from rank_eval import OUTPUT_DIR, PROJECT_DIR, top_labels

QUESTIONS_FILE = PROJECT_DIR / "questions.json"
JUDGE_MODEL = "claude-opus-5-5"   

JUDGE_SYSTEM = """你是嚴格的評分員，負責評估人資問答系統的回答品質。
你會拿到：用戶問題、標準答案、系統實際送給模型的規章條文、系統的回答。請依下列標準評分：

1. correctness（0、1 或 2）
   2：涵蓋標準答案的所有關鍵事實（結論、數字、條件），且沒有與標準答案矛盾
   1：結論大致正確，但缺少部分關鍵數字或條件；或部分正確、部分錯誤
   0：結論錯誤、與標準答案矛盾，或回答「不知道」
   用字不同但意思相同視為正確；回答提供標準答案以外的正確補充資訊不扣分。
2. faithful（true 或 false）
   回答中的每個具體主張都能在提供的條文中找到依據即為 true。
   若條文註明依據外部法規（例如勞基法）但沒有收錄其內容，回答只說明「依據該法規」為 true；
   若回答自行補充該法規的具體內容，則為 false。
3. cites（true 或 false）：回答是否註明依據的規章名稱。
4. comment：一句話說明扣分原因；滿分則寫「無」。

只輸出一個 JSON 物件，不要有其他文字：
{"correctness": 2, "faithful": true, "cites": true, "comment": "無"}"""


def judge(q: dict, context: str, answer: str) -> dict:
    user = (f"用戶問題：{q['question']}\n\n標準答案：{q['ground_truth']}\n\n"
            f"送給模型的條文：\n{context}\n\n系統的回答：\n{answer}")
    response = claude_client.messages.create(
        model=JUDGE_MODEL,
        max_tokens=512,
        system=JUDGE_SYSTEM,
        messages=[{"role": "user", "content": user}],
    )
    text = response_text(response)
    m = re.search(r"\{.*\}", text, re.DOTALL)
    if not m:
        raise ValueError(f"評分結果不是 JSON：{text[:200]}")
    return json.loads(m.group(0))


def evaluate_question(q: dict) -> dict:
    top, referenced = build_context(q["question"])
    try:
        answer = generate(q["question"], top, referenced)
    except Exception as e:                                   # 產生回答失敗：記下來，繼續下一題
        answer = f"產生回答失敗：{e}"
    context = "\n\n".join(format_chunk(c) for c in top + referenced)
    row = {
        "題號": q["id"],
        "類別": q["category"],
        "改寫類型": q.get("variant", ""),
        "問題": q["question"],
        "標準答案": q["ground_truth"],
        "回答": answer,
        f"rerank前{TOP_K_RERANKED}名": top_labels(top, n=len(top)),
        "引用帶入的條文": top_labels(referenced, n=len(referenced)),
    }
    try:
        if answer.startswith("產生回答失敗"):
            raise RuntimeError("沒有回答可以評分")
        score = judge(q, context, answer)
        row.update({
            "正確性": int(score["correctness"]),
            "完全正確": int(score["correctness"]) == 2,
            "忠實": bool(score["faithful"]),
            "註明依據": bool(score["cites"]),
            "評語": score.get("comment", ""),
        })
    except Exception as e:                                   # 評分失敗不中斷，之後人工補看
        row.update({"正確性": None, "完全正確": None, "忠實": None, "註明依據": None, "評語": f"評分失敗：{e}"})
    return row


# ---------- 匯總與輸出 ----------

METRICS = ["正確性", "完全正確", "忠實", "註明依據"]

def build_summary(df: pd.DataFrame) -> pd.DataFrame:
    scored = df.dropna(subset=["正確性"]).astype({c: float for c in METRICS})
    parts = []
    overall = scored[METRICS].mean().to_frame().T
    overall.insert(0, "題數", len(scored))
    overall.insert(0, "分組", "整體")
    parts.append(overall)
    for col in ["類別", "改寫類型"]:
        if scored[col].nunique() > 1 or (col == "改寫類型" and scored[col].ne("").any()):
            g = scored.groupby(col, sort=False)
            part = g[METRICS].mean()
            part.insert(0, "題數", g.size())
            part = part.reset_index().rename(columns={col: "分組"})
            part["分組"] = f"{col}：" + part["分組"].astype(str)
            parts.append(part)
    summary = pd.concat(parts, ignore_index=True)
    return summary.rename(columns={"正確性": "平均正確性(0-2)", "完全正確": "完全正確比例",
                                   "忠實": "忠實比例", "註明依據": "註明依據比例"})


def format_workbook(writer: pd.ExcelWriter, sheets: dict[str, pd.DataFrame]) -> None:
    from openpyxl.styles import Alignment, Font

    wide = {"問題": 30, "標準答案": 40, "回答": 70, "評語": 40}
    for name, df in sheets.items():
        ws = writer.sheets[name]
        ws.freeze_panes = "B2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for idx, col in enumerate(df.columns, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            longest = max([len(str(col))] + [len(str(v).split("\n")[0]) for v in df[col]])
            ws.column_dimensions[letter].width = wide.get(col, min(max(10, longest * 1.6), 35))
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float) and col.endswith("比例"):
                    cell.number_format = "0%"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.00"


def main() -> None:
    questions = json.loads(QUESTIONS_FILE.read_text(encoding="utf-8"))

    rows = []
    for q in questions:
        print(f"回答並評分 {q['id']} …")
        rows.append(evaluate_question(q))

    df = pd.DataFrame(rows)
    df_summary = build_summary(df)
    df_settings = pd.DataFrame([
        {"參數": "評估時間", "值": datetime.now().strftime("%Y-%m-%d %H:%M")},
        {"參數": "題目檔", "值": QUESTIONS_FILE.name},
        {"參數": "題數", "值": len(questions)},
        {"參數": "TOP_K_RETRIEVED", "值": TOP_K_RETRIEVED},
        {"參數": "送給 LLM 的 rerank 名次數", "值": TOP_K_RERANKED},
        {"參數": "引用展開最多補幾條", "值": MAX_REF_ADD},
        {"參數": "評分模型", "值": JUDGE_MODEL},
    ])

    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / f"answer_eval_{datetime.now():%Y%m%d_%H%M}.xlsx"
    sheets = {"摘要": df_summary, "回答": df, "設定": df_settings}
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for name, d in sheets.items():
            d.to_excel(writer, sheet_name=name, index=False)
        format_workbook(writer, sheets)

    overall = df_summary.iloc[0]
    print("\n===== 整體結果 =====")
    print(f"平均正確性: {overall['平均正確性(0-2)']:.2f} / 2")
    print(f"完全正確: {overall['完全正確比例']:.0%}   忠實: {overall['忠實比例']:.0%}   註明依據: {overall['註明依據比例']:.0%}")

    flagged = df[(df["正確性"] != 2) | (df["忠實"] == False)]
    if not flagged.empty:
        print("\n需要人工檢查的題目：")
        for _, r in flagged.iterrows():
            print(f"  {r['題號']}  正確性 {r['正確性']}  忠實 {r['忠實']}  {r['評語']}")
    print(f"\n已輸出：{out_path}")


if __name__ == "__main__":
    main()
